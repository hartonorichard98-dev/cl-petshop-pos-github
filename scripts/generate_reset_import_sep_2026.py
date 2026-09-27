#!/usr/bin/env python3
import hashlib
import json
import re
import uuid
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path("/Users/macmii/Downloads/(ADM) 2026's OFFLINE - Q3.xlsx")
OLD_MIGRATION = ROOT / "supabase/migrations/20260921154720_import_historical_sales_sep_2026.sql"
PRICE_LIST = ROOT / "price-list-data.js"
OUTPUT = ROOT / "supabase/migrations/20260927213000_reset_and_import_sales_through_sep_28.sql"
STORE_ID = "24a3a058-991b-4b64-b1d2-912910842d34"
IMPORT_REASON = "Historical Excel import 2026-07-01 to 2026-09-28"
NAMESPACE = uuid.UUID("8729a07b-1787-4aba-a86e-dc5828492fc8")
JAKARTA = timezone(timedelta(hours=7))


def normalize(value):
    text = str(value or "").strip().lower()
    text = text.replace("×", "x")
    return re.sub(r"\s+", " ", text)


def load_price_list():
    source = PRICE_LIST.read_text()
    match = re.search(r"window\.CL_PRICE_LIST=(\[.*\]);\s*$", source, re.S)
    if not match:
        raise RuntimeError("price-list-data.js tidak dapat dibaca")
    products = json.loads(match.group(1))
    by_barcode = {str(product.get("barcode") or "").strip(): product for product in products if product.get("barcode")}
    by_name = defaultdict(list)
    for product in products:
        for value in (product.get("name"), product.get("baseName")):
            if value:
                by_name[normalize(value)].append(product)
    return products, by_barcode, by_name


def load_historical_costs():
    source = OLD_MIGRATION.read_text()
    match = re.search(r"\$cl_import\$(\[.*\])\$cl_import\$::jsonb", source, re.S)
    if not match:
        raise RuntimeError("Payload migration historis tidak ditemukan")
    previous_sales = json.loads(match.group(1))
    cost_map = defaultdict(list)
    for sale in previous_sales:
        for item in sale["items"]:
            quantity = int(round(float(item.get("qty") or 0)))
            if quantity <= 0:
                continue
            unit_cost = float(item.get("lineCost") or 0) / quantity
            unit_sell = float(item.get("lineTotal") or 0) / quantity
            keys = {normalize(item.get("name")), normalize(item.get("baseName"))}
            for key in keys:
                if key:
                    cost_map[(key, round(unit_sell, 4))].append(unit_cost)
    return {key: sum(values) / len(values) for key, values in cost_map.items()}


def select_product(name, barcode, unit_sell, by_barcode, by_name):
    barcode_key = str(barcode or "").strip().removesuffix(".0")
    if barcode_key and barcode_key in by_barcode:
        return by_barcode[barcode_key]
    candidates = by_name.get(normalize(name), [])
    if candidates:
        return min(candidates, key=lambda product: abs(float(product.get("sell") or 0) - unit_sell))

    stripped_name = normalize(re.sub(r"^\([^)]*\)\s*", "", str(name or "")))
    candidates = by_name.get(stripped_name, [])
    if candidates:
        return min(candidates, key=lambda product: abs(float(product.get("sell") or 0) - unit_sell))

    family = stripped_name.split(" - ", 1)[0]
    family_candidates = []
    for candidate_name, products in by_name.items():
        if candidate_name.split(" - ", 1)[0] == family:
            family_candidates.extend(products)
    if family_candidates:
        return min(family_candidates, key=lambda product: abs(float(product.get("sell") or 0) - unit_sell))
    return None


def parse_payment(label):
    value = normalize(label)
    if "split" in value:
        return "split"
    if "qris" in value:
        return "qris"
    if "trf" in value or "transfer" in value:
        return "transfer"
    if "debit" in value:
        return "debit"
    return "cash"


def parse_split(note, total):
    aliases = {"c": "cash", "q": "qris", "t": "transfer", "d": "debit"}
    parts = []
    for amount_a, method_a, method_b, amount_b in re.findall(r"(?:(\d+(?:[.,]\d+)?)\s*([cqtd])|([cqtd])\s*(\d+(?:[.,]\d+)?))", normalize(note)):
        method = method_a or method_b
        raw_amount = amount_a or amount_b
        amount = int(round(float(raw_amount.replace(",", ".")) * 1000))
        parts.append({"method": aliases[method], "amount": amount})
    if len(parts) != 2:
        raise ValueError(f"Split payment tidak valid: note={note!r}, total={total}, parsed={parts}")
    if sum(part["amount"] for part in parts) != total:
        # Excel notes occasionally round amounts in thousands. Keep first leg,
        # then reconcile second leg to the authoritative transaction total.
        parts[1]["amount"] = total - parts[0]["amount"]
        if parts[1]["amount"] <= 0:
            raise ValueError(f"Split payment tidak valid: note={note!r}, total={total}, parsed={parts}")
    if parts[0]["method"] == parts[1]["method"]:
        raise ValueError(f"Split payment memakai metode sama: {note!r}")
    return parts


def build_sales():
    _, by_barcode, by_name = load_price_list()
    historical_costs = load_historical_costs()
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=False)
    raw_transactions = []

    sales_sheets = (
        "JUL 1", "JUL 2", "JUL 3", "JUL 4", "JUL 5",
        "AUG 1", "AUG 2", "AUG 3", "AUG 4", "AUG 5",
        "SEP 1", "SEP 2", "SEP 3", "SEP 4", "SEP 5",
    )
    for sheet_name in sales_sheets:
        worksheet = workbook[sheet_name]
        current = None
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            date_value, transaction_number, product_name, barcode, fallback_name, price, quantity, total, payment, note = row[:10]
            name = product_name or fallback_name
            if not isinstance(date_value, datetime) or not name or not total or float(total) <= 0:
                continue
            has_transaction_number = transaction_number is not None and str(transaction_number).strip() != ""
            if has_transaction_number:
                current = {
                    "sheet": sheet_name,
                    "date": date_value.date().isoformat(),
                    "trx": str(int(transaction_number)) if float(transaction_number).is_integer() else str(transaction_number),
                    "paymentLabel": payment,
                    "paymentNote": note,
                    "rows": [],
                    "sourceRow": row_number,
                }
                raw_transactions.append(current)
            if current is None:
                raise ValueError(f"Baris tanpa nomor transaksi awal: {sheet_name}!{row_number}")
            if current["date"] != date_value.date().isoformat():
                raise ValueError(f"Tanggal berubah di tengah transaksi: {sheet_name}!{row_number}")
            current["rows"].append({
                "name": str(name).strip(),
                "barcode": str(barcode or "").strip().removesuffix(".0"),
                "price": float(price or 0),
                "quantity": int(round(float(quantity or 0))),
                "total": int(round(float(total))),
                "row": row_number,
            })

    sales = []
    unmatched = []
    global_index = 0
    for raw in raw_transactions:
        global_index += 1
        items = []
        for source in raw["rows"]:
            quantity = source["quantity"]
            if quantity <= 0:
                raise ValueError(f"Qty tidak valid: {raw['sheet']}!{source['row']}")
            unit_sell = source["total"] / quantity
            product = select_product(source["name"], source["barcode"], unit_sell, by_barcode, by_name)
            if not product:
                unmatched.append((raw["sheet"], source["row"], source["name"], source["barcode"], unit_sell))
                product = {
                    "id": None,
                    "sku": "",
                    "barcode": source["barcode"],
                    "name": source["name"],
                    "baseName": re.sub(r"^\([^)]*\)\s*", "", source["name"]).strip(),
                    "sell": unit_sell,
                    "cost": 0,
                    "active": False,
                    "trackStock": False,
                    "isTier": source["name"].startswith("("),
                    "tierMin": 1,
                }
            cost_key_candidates = (
                (normalize(source["name"]), round(unit_sell, 4)),
                (normalize(product.get("name")), round(unit_sell, 4)),
                (normalize(product.get("baseName")), round(unit_sell, 4)),
            )
            unit_cost = next((historical_costs[key] for key in cost_key_candidates if key in historical_costs), float(product.get("cost") or 0))
            line_cost = round(unit_cost * quantity, 2)
            tier_min = int(product.get("tierMin") or 1)
            items.append({
                "id": product.get("id"),
                "sku": product.get("sku") or "",
                "barcode": product.get("barcode") or source["barcode"],
                "name": source["name"],
                "baseName": product.get("baseName") or source["name"],
                "sell": unit_sell,
                "cost": unit_cost,
                "qty": quantity,
                "lineCost": line_cost,
                "lineTotal": source["total"],
                "active": bool(product.get("active", True)),
                "trackStock": bool(product.get("trackStock", False)),
                "isTier": bool(product.get("isTier", False)),
                "tierMin": tier_min,
                "minQty": tier_min,
                "groupKey": product.get("baseName") or source["name"],
                "image": product.get("image") or "",
                "needsPrice": bool(product.get("needsPrice", False)),
            })

        total = sum(item["lineTotal"] for item in items)
        payment_method = parse_payment(raw["paymentLabel"])
        payment_breakdown = parse_split(raw["paymentNote"], total) if payment_method == "split" else [{"method": payment_method, "amount": total}]
        cash_received = sum(part["amount"] for part in payment_breakdown if part["method"] == "cash")
        source_key = f"{raw['sheet']}|{raw['date']}|{raw['trx']}|{raw['sourceRow']}|{total}"
        sale_id = str(uuid.uuid5(NAMESPACE, source_key))
        created_at = datetime.fromisoformat(raw["date"]).replace(tzinfo=JAKARTA, hour=12) + timedelta(seconds=global_index)
        sales.append({
            "date": raw["date"],
            "trx": raw["trx"],
            "items": items,
            "total": total,
            "paymentMethod": payment_method,
            "paymentBreakdown": payment_breakdown,
            "cashReceived": cash_received,
            "changeDue": 0,
            "costTotal": int(round(sum(float(item["lineCost"]) for item in items))),
            "signature": hashlib.sha256(source_key.encode()).hexdigest(),
            "id": sale_id,
            "receipt": f"RESET-{sale_id[:12]}",
            "cashier": "Historical Excel import",
            "status": "completed",
            "reason": IMPORT_REASON,
            "paymentEditHistory": ([{"action": "historical_import", "note": f"Excel split note: {raw['paymentNote']}"}] if payment_method == "split" and raw["paymentNote"] else []),
            "createdAt": created_at.isoformat(),
        })

    return sales, unmatched


def sql_literal_json(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def write_migration(sales):
    expected_count = len(sales)
    expected_total = sum(sale["total"] for sale in sales)
    min_date = min(sale["date"] for sale in sales)
    max_date = max(sale["date"] for sale in sales)
    payload = sql_literal_json(sales)
    sql = f"""begin;

do $reset_guard$
begin
  if not exists (select 1 from public.stores where id = '{STORE_ID}'::uuid) then
    raise exception 'Target store not found';
  end if;
end
$reset_guard$;

delete from public.inventory_movement_change_requests where store_id = '{STORE_ID}'::uuid;
delete from public.inventory_movements where store_id = '{STORE_ID}'::uuid;
delete from public.sale_deletions where store_id = '{STORE_ID}'::uuid;
delete from public.audit_logs where store_id = '{STORE_ID}'::uuid;
delete from public.cash_closings where store_id = '{STORE_ID}'::uuid;
delete from public.receivings where store_id = '{STORE_ID}'::uuid;
delete from public.expenses where store_id = '{STORE_ID}'::uuid;
delete from public.sales where store_id = '{STORE_ID}'::uuid;
delete from public.pos_receipt_sequences where store_id = '{STORE_ID}'::uuid;

update public.products
set stock = 75,
    track_stock = true,
    updated_at = now()
where store_id = '{STORE_ID}'::uuid;

with payload as (
  select jsonb_array_elements($cl_reset_import${payload}$cl_reset_import$::jsonb) as sale
), inserted as (
  insert into public.sales (
    id, store_id, receipt_number, cashier_id, cashier_name, items, total,
    cost_total, payment_method, cash_received, change_due, status,
    correction_reason, business_date, created_at, updated_at,
    payment_breakdown, payment_edit_history
  )
  select
    (sale->>'id')::uuid,
    '{STORE_ID}'::uuid,
    sale->>'receipt',
    null,
    sale->>'cashier',
    sale->'items',
    (sale->>'total')::bigint,
    (sale->>'costTotal')::bigint,
    sale->>'paymentMethod',
    (sale->>'cashReceived')::bigint,
    (sale->>'changeDue')::bigint,
    'completed'::public.sale_status,
    sale->>'reason',
    (sale->>'date')::date,
    (sale->>'createdAt')::timestamptz,
    (sale->>'createdAt')::timestamptz,
    sale->'paymentBreakdown',
    '[]'::jsonb
  from payload
  order by (sale->>'createdAt')::timestamptz
  returning id
)
select count(*) as inserted_sales from inserted;

do $validation$
declare
  actual_sales bigint;
  actual_total bigint;
  actual_min_date date;
  actual_max_date date;
  wrong_stock bigint;
  trx_85_count bigint;
  sep_28_non_cash bigint;
begin
  select count(*), coalesce(sum(total), 0), min(business_date), max(business_date)
  into actual_sales, actual_total, actual_min_date, actual_max_date
  from public.sales
  where store_id = '{STORE_ID}'::uuid;

  if actual_sales <> {expected_count} or actual_total <> {expected_total}
     or actual_min_date <> date '{min_date}' or actual_max_date <> date '{max_date}' then
    raise exception 'Sales validation failed: count %, total %, range % to %', actual_sales, actual_total, actual_min_date, actual_max_date;
  end if;

  select count(*) into wrong_stock
  from public.products
  where store_id = '{STORE_ID}'::uuid and (stock <> 75 or track_stock is not true);
  if wrong_stock <> 0 then
    raise exception 'Product stock validation failed: % rows', wrong_stock;
  end if;

  select count(*) into trx_85_count
  from public.sales
  where store_id = '{STORE_ID}'::uuid
    and business_date = date '2026-09-26'
    and payment_method = 'transfer'
    and total = 294000
    and items @> '[{{"name":"cat choize kitten 1 kg - salmon","qty":2}}]'::jsonb
    and items @> '[{{"name":"cat choize kitten 1 kg - tuna","qty":2}}]'::jsonb
    and items @> '[{{"name":"meo kaleng 400 gr - kitten tuna sardines","qty":10}}]'::jsonb;
  if trx_85_count <> 1 then
    raise exception 'Transaction 85 validation failed: % matches', trx_85_count;
  end if;

  select count(*) into sep_28_non_cash
  from public.sales
  where store_id = '{STORE_ID}'::uuid
    and business_date = date '2026-09-28'
    and payment_method <> 'cash';
  if sep_28_non_cash <> 0 then
    raise exception 'September 28 payment validation failed: % non-cash rows', sep_28_non_cash;
  end if;
end
$validation$;

select
  count(*) as sales_count,
  sum(total) as sales_total,
  min(business_date) as first_business_date,
  max(business_date) as last_business_date,
  count(*) filter (where payment_method = 'cash') as cash_sales,
  count(*) filter (where payment_method = 'qris') as qris_sales,
  count(*) filter (where payment_method = 'transfer') as transfer_sales,
  count(*) filter (where payment_method = 'debit') as debit_sales,
  count(*) filter (where payment_method = 'split') as split_sales
from public.sales
where store_id = '{STORE_ID}'::uuid;

select count(*) as product_count, min(stock) as min_stock, max(stock) as max_stock
from public.products
where store_id = '{STORE_ID}'::uuid;

commit;
"""
    OUTPUT.write_text(sql)
    return expected_count, expected_total


if __name__ == "__main__":
    sales, unmatched = build_sales()
    count, total = write_migration(sales)
    payment_counts = defaultdict(int)
    for sale in sales:
        payment_counts[sale["paymentMethod"]] += 1
    print(json.dumps({
        "output": str(OUTPUT),
        "sales": count,
        "total": total,
        "dateRange": [min(sale["date"] for sale in sales), max(sale["date"] for sale in sales)],
        "payments": dict(sorted(payment_counts.items())),
        "unmatchedProducts": unmatched,
    }, ensure_ascii=False, indent=2))
