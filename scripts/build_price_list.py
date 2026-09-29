import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl


def barcode_value(value: object) -> str:
    if value in (None, ''):
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def main() -> None:
    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if 'DATABASE Q4' not in workbook.sheetnames:
        raise ValueError("Sheet 'DATABASE Q4' tidak ditemukan")
    sheet = workbook['DATABASE Q4']
    groups = OrderedDict()
    source_product_id = 0

    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number == 1:
            continue

        name = str(row[0] or '').strip()
        if not name:
            continue
        source_product_id += 1

        sell_price = row[1] if len(row) > 1 and isinstance(row[1], (int, float)) else 0
        cost_price = row[2] if len(row) > 2 and isinstance(row[2], (int, float)) else 0
        barcode = barcode_value(row[3]) if len(row) > 3 else ''
        has_price = sell_price > 0
        tier_match = re.match(r'^\((\d+)\+\)\s*(.+)$', name, re.IGNORECASE)
        unit_match = re.match(r'^\(SATUAN\)\s*(.+)$', name, re.IGNORECASE)
        tier_minimum = int(tier_match.group(1)) if tier_match else 1
        base_name = tier_match.group(2).strip() if tier_match else unit_match.group(1).strip() if unit_match else name

        group = groups.setdefault(base_name.casefold(), {'name': base_name, 'rows': []})
        group['rows'].append(
            {
                'sourceName': name,
                'sourceId': source_product_id,
                'minQty': tier_minimum,
                'sell': round(float(sell_price), 2),
                'cost': round(float(cost_price), 2) if has_price else 0,
                'barcode': barcode,
                'hasPrice': has_price,
            }
        )

    products = []
    for group in groups.values():
        priced_rows = sorted(
            (row for row in group['rows'] if row['hasPrice']),
            key=lambda row: row['minQty'],
        )
        primary = next((row for row in priced_rows if row['minQty'] == 1), priced_rows[0] if priced_rows else group['rows'][0])
        tiers_by_minimum = OrderedDict()
        for row in priced_rows:
            if row['minQty'] in tiers_by_minimum:
                raise ValueError(f"Tangga harga {row['minQty']}+ duplikat untuk {group['name']}")
            tiers_by_minimum[row['minQty']] = {
                'minQty': row['minQty'],
                'sell': row['sell'],
                'cost': row['cost'],
            }
        if group['name'].casefold() == 'liebao 15 gr' and 1 not in tiers_by_minimum:
            reference = tiers_by_minimum.get(5) or primary
            tiers_by_minimum[1] = {'minQty': 1, 'sell': 1000, 'cost': reference['cost']}
            tiers_by_minimum = OrderedDict(sorted(tiers_by_minimum.items()))
        product_id = min(row['sourceId'] for row in group['rows'])
        barcode = primary['barcode'] or next((row['barcode'] for row in group['rows'] if row['barcode']), '')
        products.append(
            {
                'id': product_id,
                'sku': f'CL-{product_id:04d}',
                'barcode': barcode,
                'image': '',
                'name': group['name'],
                'baseName': group['name'],
                'sell': primary['sell'] if primary['hasPrice'] else 0,
                'cost': primary['cost'] if primary['hasPrice'] else 0,
                'tiers': list(tiers_by_minimum.values()),
                'stock': 0,
                'trackStock': False,
                'active': bool(priced_rows),
                'needsPrice': not priced_rows,
            }
        )

    generated_by_name = {product['baseName'].casefold(): product for product in products}
    for group_name, group in groups.items():
        product = generated_by_name[group_name]
        generated_tiers = {tier['minQty']: tier for tier in product['tiers']}
        for row in (item for item in group['rows'] if item['hasPrice']):
            generated = generated_tiers.get(row['minQty'])
            if not generated:
                raise ValueError(f"Tangga harga {row['minQty']}+ hilang untuk {group['name']}")
            if generated['sell'] != row['sell'] or generated['cost'] != row['cost']:
                raise ValueError(f"Harga berubah saat generate untuk {group['name']} tier {row['minQty']}+")

    payload = json.dumps(products, ensure_ascii=False, separators=(',', ':'))
    destination.write_text(
        "window.CL_PRICE_LIST_VERSION='2026-09-29-database-q4-v3';\n"
        f'window.CL_PRICE_LIST={payload};\n',
        encoding='utf-8',
    )

    missing = [product for product in products if product['needsPrice']]
    tier_count = sum(len(product['tiers']) for product in products)
    print(f'products={len(products)} tiers={tier_count} missing_price={len(missing)}')
    for product in missing:
        print(f"missing: {product['sku']} {product['name']}")


if __name__ == '__main__':
    main()
